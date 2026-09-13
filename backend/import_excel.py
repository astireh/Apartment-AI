import sqlite3
from openpyxl import load_workbook


EXCEL_FILE = "apartments.xlsx"
DATABASE_FILE = "apartments.db"


def import_excel_to_database():
    """Импортирует и обновляет данные из Excel в SQLite."""

    print()
    print("Начинаю импорт данных из Excel...")
    print()

    # Загружаем Excel
    workbook = load_workbook(
        EXCEL_FILE,
        data_only=True
    )

    sheet_complexes = workbook["ЖК"]
    sheet_apartments = workbook["Квартиры"]

    # Подключаемся к SQLite
    connection = sqlite3.connect(
        DATABASE_FILE
    )

    cursor = connection.cursor()

    # -------------------------------------------------
    # Создание таблиц
    # -------------------------------------------------

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS residential_complexes (
            id INTEGER PRIMARY KEY,
            name TEXT NOT NULL,
            district TEXT,
            housing_class TEXT,
            completion_year INTEGER,
            description TEXT,
            infrastructure TEXT,
            parking TEXT,
            image TEXT
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS apartments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            complex_id INTEGER NOT NULL,
            rooms INTEGER NOT NULL,
            area REAL NOT NULL,
            price INTEGER NOT NULL,
            floor INTEGER NOT NULL,
            finishing TEXT,
            balcony TEXT,
            image TEXT,
            booking_status TEXT NOT NULL DEFAULT 'available',
            booking_type TEXT,
            booking_until TEXT,
            FOREIGN KEY (complex_id)
                REFERENCES residential_complexes(id)
        )
    """)

    # -------------------------------------------------
    # Проверяем наличие image у старой таблицы ЖК
    # -------------------------------------------------

    cursor.execute(
        "PRAGMA table_info(residential_complexes)"
    )

    complex_columns = [
        row[1]
        for row in cursor.fetchall()
    ]

    if "image" not in complex_columns:
        cursor.execute("""
            ALTER TABLE residential_complexes
            ADD COLUMN image TEXT
        """)

    # -------------------------------------------------
    # Импорт ЖК
    # -------------------------------------------------

    complex_added = 0
    complex_updated = 0

    for row in sheet_complexes.iter_rows(
        min_row=2,
        values_only=True
    ):

        if not row[0]:
            continue

        (
            complex_id,
            name,
            district,
            housing_class,
            completion_year,
            description,
            infrastructure,
            parking,
            image
        ) = row

        # Проверяем, существует ли ЖК
        cursor.execute("""
            SELECT id
            FROM residential_complexes
            WHERE id = ?
        """, (complex_id,))

        existing_complex = cursor.fetchone()

        if existing_complex:
            # Обновляем существующий ЖК
            cursor.execute("""
                UPDATE residential_complexes
                SET
                    name = ?,
                    district = ?,
                    housing_class = ?,
                    completion_year = ?,
                    description = ?,
                    infrastructure = ?,
                    parking = ?,
                    image = ?
                WHERE id = ?
            """, (
                name,
                district,
                housing_class,
                completion_year,
                description,
                infrastructure,
                parking,
                image,
                complex_id
            ))

            complex_updated += 1

        else:
            # Добавляем новый ЖК
            cursor.execute("""
                INSERT INTO residential_complexes (
                    id,
                    name,
                    district,
                    housing_class,
                    completion_year,
                    description,
                    infrastructure,
                    parking,
                    image
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                complex_id,
                name,
                district,
                housing_class,
                completion_year,
                description,
                infrastructure,
                parking,
                image
            ))

            complex_added += 1

    # -------------------------------------------------
    # Импорт квартир
    # -------------------------------------------------

    apartment_added = 0
    apartment_updated = 0

    for row in sheet_apartments.iter_rows(
        min_row=2,
        values_only=True
    ):

        if not row[0]:
            continue

        (
            apartment_id,
            complex_id,
            rooms,
            area,
            price,
            floor,
            finishing,
            balcony,
            image
        ) = row

        # Проверяем, существует ли квартира
        cursor.execute("""
            SELECT id
            FROM apartments
            WHERE id = ?
        """, (apartment_id,))

        existing_apartment = cursor.fetchone()

        if existing_apartment:
            # Обновляем только данные квартиры.
            # Данные бронирования НЕ трогаем.
            cursor.execute("""
                UPDATE apartments
                SET
                    complex_id = ?,
                    rooms = ?,
                    area = ?,
                    price = ?,
                    floor = ?,
                    finishing = ?,
                    balcony = ?,
                    image = ?
                WHERE id = ?
            """, (
                complex_id,
                rooms,
                area,
                price,
                floor,
                finishing,
                balcony,
                image,
                apartment_id
            ))

            apartment_updated += 1

        else:
            # Добавляем новую квартиру.
            # Для неё booking_status автоматически будет "available".
            cursor.execute("""
                INSERT INTO apartments (
                    id,
                    complex_id,
                    rooms,
                    area,
                    price,
                    floor,
                    finishing,
                    balcony,
                    image
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                apartment_id,
                complex_id,
                rooms,
                area,
                price,
                floor,
                finishing,
                balcony,
                image
            ))

            apartment_added += 1

    # -------------------------------------------------
    # Сохраняем изменения
    # -------------------------------------------------

    connection.commit()

    # Закрываем соединение
    connection.close()

    # -------------------------------------------------
    # Результат
    # -------------------------------------------------

    print(
        f"ЖК добавлено: {complex_added}"
    )

    print(
        f"ЖК обновлено: {complex_updated}"
    )

    print(
        f"Квартир добавлено: {apartment_added}"
    )

    print(
        f"Квартир обновлено: {apartment_updated}"
    )

    print()
    print(
        "Импорт Excel → SQLite завершён успешно."
    )
    print()


if __name__ == "__main__":
    import_excel_to_database()